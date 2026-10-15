"""Isolated smoke test for save_jobs/export_report without network/API calls."""
import ast, json, logging, sqlite3, tempfile
from pathlib import Path
from datetime import datetime, timezone
from typing import Any
from urllib.parse import urlparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SOURCE = (Path(__file__).resolve().parent.parent / 'main.py').read_text(encoding='utf-8')
TREE = ast.parse(SOURCE)
WANTED = {'save_jobs', 'export_report'}
segments = []
for node in TREE.body:
    if isinstance(node, ast.FunctionDef) and node.name in WANTED:
        segments.append(ast.get_source_segment(SOURCE, node))
ns = {
    'Any': Any, 'JobSearchState': dict, 'json': json, 'sqlite3': sqlite3, 'datetime': datetime,
    'timezone': timezone, 'Path': Path, 'Workbook': Workbook, 'Font': Font,
    'PatternFill': PatternFill, 'Alignment': Alignment,
    'get_column_letter': get_column_letter, 'logger': logging.getLogger('test'),
    'normalize_url': lambda u: str(u or '').strip().rstrip('/').split('#')[0],
    '_fallback_role_category': lambda title: 'Network & NOC Engineering' if 'noc' in str(title).lower() else 'Other IT Roles',
    'STRONG_FIT_MIN_SCORE': 75,
}
exec('\n\n'.join(segments), ns)
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    ns['DB_PATH'] = tmp/'jobs.db'
    ns['OUTPUT_DIR'] = tmp/'output'
    ns['OUTPUT_DIR'].mkdir()
    ns['CV_PATH'] = Path('my_cv.pdf')
    with sqlite3.connect(ns['DB_PATH']) as con:
        con.execute('''CREATE TABLE jobs (
          url TEXT PRIMARY KEY, title TEXT, company TEXT, location TEXT, region TEXT,
          description TEXT, fit_score INTEGER, reason TEXT, missing_skills TEXT,
          recommended INTEGER, first_seen TEXT, last_seen TEXT,
          email_sent INTEGER NOT NULL DEFAULT 0, category_email_sent INTEGER NOT NULL DEFAULT 0,
          job_category TEXT NOT NULL DEFAULT '', verification_status TEXT NOT NULL DEFAULT 'unverified',
          verification_reason TEXT NOT NULL DEFAULT '')''')
    state = {
      'unique_jobs': [
        {'url':'https://example.com/jobs/abc','title':'NOC Engineer','description':'NOC role short snippet','company':'Example','location':'Brno','region':'Europe'},
        {'url':'https://example.com/careers','title':'Careers page','description':'generic page'},
        {'url':'https://example.com/course/devops-course','title':'DevOps Course','description':'course'}
      ],
      'review_required_jobs': [
        {'url':'https://example.com/jobs/abc','title':'NOC Engineer','verification_status':'review_required','verification_reason':'403 blocked'}
      ],
      'excluded_pages': [
        {'url':'https://example.com/course/devops-course','title':'DevOps Course','verification_status':'not_job','verification_reason':'course page'}
      ],
      'verified_jobs': [],
      'assessed_jobs': [],
    }
    ns['save_jobs'](state)
    ns['export_report'](state)
    with sqlite3.connect(ns['DB_PATH']) as con:
       rows = con.execute('SELECT url, verification_status FROM jobs ORDER BY url').fetchall()
    wb = load_workbook(ns['OUTPUT_DIR']/'job_history_my_cv.xlsx', read_only=False, data_only=True)
    names = wb.sheetnames
    sheet_counts = {name: wb[name].max_row for name in names}
    wb.close()
    assert len(rows) == 3, rows
    assert set(names) == {'All Discovered Leads','Verified & Scored','Needs Manual Review','Excluded Non-Jobs'}, names
    assert sheet_counts['All Discovered Leads'] == 4, sheet_counts
    assert sheet_counts['Needs Manual Review'] >= 2, sheet_counts
    assert sheet_counts['Excluded Non-Jobs'] == 2, sheet_counts
    all_sheet = wb['All Discovered Leads']
    headers = [cell.value for cell in all_sheet[1]]
    url_col = headers.index('Job URL') + 1
    stored_urls = [all_sheet.cell(row=r, column=url_col).value for r in range(2, all_sheet.max_row + 1)]
    assert 'https://example.com/jobs/abc' in stored_urls, stored_urls
    url_row = next(r for r in range(2, all_sheet.max_row + 1) if all_sheet.cell(row=r, column=url_col).value == 'https://example.com/jobs/abc')
    assert all_sheet.cell(row=url_row, column=url_col).hyperlink is not None, 'Job URL should be clickable in Excel'
    print('PASS: database persisted 3 URLs, and workbook has review/excluded rows and all 4 sheets:', sheet_counts)
