"""Prepare job-specific application documents and route jobs to authorized adapters.

No website is automatically submitted to unless a separately implemented,
platform-authorized adapter is registered in application_adapters.py.
"""
from __future__ import annotations

import json
import logging
import hashlib
from html import escape
import os
import re
import sqlite3
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, Field
from dotenv import load_dotenv

load_dotenv()

from email_notifier import send_application_report_email

logger = logging.getLogger(__name__)

APPLICATION_MIN_FIT_SCORE = int(os.getenv("APPLICATION_MIN_FIT_SCORE", "80"))
MAX_APPLICATION_PACKAGES_PER_RUN = max(1, int(os.getenv("MAX_APPLICATION_PACKAGES_PER_RUN", "10")))
AUTO_APPLY_ENABLED = os.getenv("AUTO_APPLY_ENABLED", "false").strip().lower() in {"1", "true", "yes", "on"}


class ExperienceItem(BaseModel):
    role: str = Field(description="Role/title exactly supported by the original CV")
    company: str = Field(description="Employer exactly supported by the original CV")
    dates: str = Field(description="Employment dates exactly supported by the original CV, or empty")
    bullets: list[str] = Field(description="Tailored bullets grounded only in the original CV; do not invent metrics")


class ProjectItem(BaseModel):
    name: str = Field(description="Project name from the original CV")
    technologies: list[str] = Field(description="Technologies explicitly supported by the original CV")
    bullets: list[str] = Field(description="Relevant project bullets grounded only in the original CV")


class EducationItem(BaseModel):
    degree: str
    institution: str
    dates: str = ""
    details: list[str] = Field(default_factory=list)


class TailoredApplication(BaseModel):
    candidate_name: str = Field(description="Candidate name as shown in the CV; do not guess")
    contact_line: str = Field(description="Contact details copied from the CV if present; otherwise empty")
    headline: str = Field(description="Targeted professional headline without adding qualifications")
    summary: str = Field(description="Tailored professional summary supported by the original CV")
    skills: list[str] = Field(description="Relevant skills that are explicitly evidenced in the CV")
    experience: list[ExperienceItem] = Field(description="Relevant work experience from the original CV; preserve employer and dates")
    projects: list[ProjectItem] = Field(description="Relevant projects from the original CV; don't invent projects")
    education: list[EducationItem] = Field(description="Education from the original CV")
    certifications: list[str] = Field(description="Certifications/courses explicitly stated in the original CV")
    cover_letter: str = Field(description="A complete, professional cover letter tailored to this specific job")


def _safe_slug(value: str, max_len: int = 65) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip()).strip("._-")
    return (value or "job")[:max_len]


def _host(url: str) -> str:
    try:
        return (urlparse(url).hostname or "unknown").lower().removeprefix("www.")
    except Exception:
        return "unknown"


def ensure_application_table(db_path: Path) -> None:
    with sqlite3.connect(db_path) as connection:
        connection.execute(
            """CREATE TABLE IF NOT EXISTS applications (
                url TEXT PRIMARY KEY,
                title TEXT NOT NULL DEFAULT '',
                company TEXT NOT NULL DEFAULT '',
                location TEXT NOT NULL DEFAULT '',
                fit_score INTEGER NOT NULL DEFAULT 0,
                platform TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'pending_preparation',
                cv_path TEXT NOT NULL DEFAULT '',
                cover_letter_path TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                submitted_at TEXT NOT NULL DEFAULT '',
                confirmation TEXT NOT NULL DEFAULT '',
                report_sent INTEGER NOT NULL DEFAULT 0
            )"""
        )
        columns = {row[1] for row in connection.execute("PRAGMA table_info(applications)")}
        if "report_sent" not in columns:
            connection.execute("ALTER TABLE applications ADD COLUMN report_sent INTEGER NOT NULL DEFAULT 0")



def _is_listing_page(job: dict[str, Any]) -> bool:
    """Reject broad job-search pages and company career homepages, not job details."""
    title = " ".join(str(job.get("title") or "").lower().split())
    url = str(job.get("url") or "").strip()
    if not url:
        return True
    try:
        parsed = urlparse(url)
        host = (parsed.hostname or "").lower().removeprefix("www.")
        path = re.sub(r"/{2,}", "/", unquote(parsed.path).lower()).rstrip("/") or "/"
        query = parse_qs(parsed.query)
        query_keys = {key.lower() for key in query}
    except Exception:
        return True

    # Known generic result pages / career homepages seen in this tracker.
    if host.endswith("remoterocketship.com") and path.startswith("/jobs/"):
        return True
    if host.endswith("naukrigulf.com") and (
        path in {"/jobs", "/search-jobs", "/job-search", "/jobs-in"}
        or re.search(r"(?:^|/)[^/]+-jobs(?:-in-[^/]+|-by-[^/]+)?$", path)
    ):
        return True
    if host.endswith("gulftalent.com") and path.startswith("/companies/"):
        return True
    if host.endswith("devopshunt.com") and path == "/":
        return True
    if host.endswith("redhat.com") and path in {"/jobs", "/en/jobs"}:
        return True
    if host.endswith("etisalat.ae") and path in {"/", "/en", "/en/index.html"}:
        return True
    if host.endswith("careers.microsoft.com") and path in {"/", "/us/en", "/us/en/"}:
        return True
    if host.endswith("jooble.org") and (
        path in {"/", "/search"} or path.startswith("/práce-") or path.startswith("/prace-")
    ):
        return True
    if host.endswith("jobs.cz") and (path == "/prace" or query_keys.intersection({"q", "q[]"})):
        return True

    # Preserve known individual-job endpoints.
    if host.endswith("indeed.com") and path == "/viewjob" and query.get("jk"):
        return False
    if host.endswith("linkedin.com") and re.match(r"^/jobs/view/\d+", path):
        return False

    if host.endswith("linkedin.com") and (
        path.startswith("/jobs/search") or path.startswith("/jobs/collections/")
        or (path == "/jobs" and query_keys.intersection({"keywords", "location", "geoid", "f_tpr"}))
    ):
        return True
    if host.endswith("indeed.com") and (
        path in {"/jobs", "/jobs/"} or path.startswith("/jobs/what-")
        or re.search(r"(?:^|/)[^/]+-jobs(?:\.html)?$", path)
    ):
        return True

    exact_listing_paths = {
        "/search", "/jobs/search", "/job/search", "/jobs/search-results",
        "/job-search", "/search-jobs", "/career-search", "/careers/search",
        "/vacancies/search", "/browse-jobs", "/find-jobs", "/job-listings",
        "/jobs/list", "/jobs/all", "/careers", "/vacancies",
    }
    if path in exact_listing_paths:
        return True
    broad_index_paths = {"/jobs", "/job", "/search", "/careers", "/vacancies", "/job-search"}
    if path in broad_index_paths and query_keys.intersection(
        {"q", "query", "keywords", "search", "fromage", "start", "page", "location", "what", "where"}
    ):
        return True
    listing_title_patterns = (
        r"\b\d[\d,]*\+?\s+(?:new\s+)?jobs?\b",
        r"\bjobs?\s+(?:in|near|available|found|matching)\b",
        r"\b(?:search|find|browse)\s+(?:for\s+)?jobs?\b",
        r"\bjob listings?\b", r"\bjob search\b", r"\ball jobs\b",
        r"\bjobs and careers\b", r"\b\d+\s+job openings\b",
        r"\b(?:search results|results for)\b",
    )
    return any(re.search(pattern, title) for pattern in listing_title_patterns)

def _build_prompt(cv_text: str, profile: dict[str, Any], job: dict[str, Any]) -> str:
    return f"""You are preparing a truthful, ATS-friendly job application.

TARGET JOB
Title: {job.get('title', '')}
Company: {job.get('company', '')}
Location: {job.get('location', '')}
URL: {job.get('url', '')}
Description:
{str(job.get('description', ''))[:9000]}

CV-DERIVED PROFILE
{json.dumps(profile, ensure_ascii=False)}

ORIGINAL CV (source of truth)
{cv_text[:35000]}

Return a tailored CV and cover letter using the required schema. Rules:
- The original CV is the only source of truth for the candidate's background.
- Never invent skills, certifications, degrees, employers, dates, projects, metrics, achievements, contact information, work authorization, or years of experience.
- Do not claim the candidate has a missing skill. Only include skills evidenced in the original CV.
- Keep employer names, job titles, education names, and dates faithful to the original.
- Reorder and rephrase relevant facts to align with this particular vacancy; do not misrepresent them.
- Include relevant projects and experience, but do not create new ones.
- If contact details are not in the CV, leave contact_line empty.
- Use a professional, concise, ATS-friendly structure. Avoid tables, graphics, icons, columns, and keyword stuffing.
- The cover letter must name the target role and company where known, explain fit using real evidence, and not claim anything unsupported. If the job description is generic or missing, write a conservative letter.
"""


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    if not text.strip():
        return
    p = document.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)


def _add_bullets(document: Document, values: list[str]) -> None:
    for value in values:
        if str(value).strip():
            document.add_paragraph(str(value).strip(), style="List Bullet")


def _write_cv_docx(data: TailoredApplication, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(35, 35, 35)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data.candidate_name or "Candidate")
    run.bold = True
    run.font.size = Pt(18)
    if data.contact_line.strip():
        p = doc.add_paragraph(data.contact_line.strip())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if data.headline.strip():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(data.headline.strip())
        r.bold = True
        r.font.size = Pt(11)

    _add_heading(doc, "Professional Summary")
    doc.add_paragraph(data.summary.strip())

    if data.skills:
        _add_heading(doc, "Technical Skills")
        doc.add_paragraph(" • ".join(dict.fromkeys(s.strip() for s in data.skills if s.strip())))

    if data.experience:
        _add_heading(doc, "Professional Experience")
        for item in data.experience:
            p = doc.add_paragraph()
            r = p.add_run(item.role.strip())
            r.bold = True
            company_dates = " | ".join(part.strip() for part in (item.company, item.dates) if part and part.strip())
            if company_dates:
                p.add_run(" — " + company_dates)
            _add_bullets(doc, item.bullets)

    if data.projects:
        _add_heading(doc, "Selected Projects")
        for item in data.projects:
            p = doc.add_paragraph()
            r = p.add_run(item.name.strip())
            r.bold = True
            if item.technologies:
                p.add_run(" | " + ", ".join(item.technologies))
            _add_bullets(doc, item.bullets)

    if data.education:
        _add_heading(doc, "Education")
        for item in data.education:
            p = doc.add_paragraph()
            r = p.add_run(item.degree.strip())
            r.bold = True
            institution_dates = " | ".join(part.strip() for part in (item.institution, item.dates) if part and part.strip())
            if institution_dates:
                p.add_run(" — " + institution_dates)
            _add_bullets(doc, item.details)

    if data.certifications:
        _add_heading(doc, "Certifications and Training")
        _add_bullets(doc, data.certifications)

    doc.save(path)


def _write_cover_letter_docx(data: TailoredApplication, job: dict[str, Any], path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    if data.candidate_name.strip():
        p = doc.add_paragraph()
        r = p.add_run(data.candidate_name.strip())
        r.bold = True
        r.font.size = Pt(14)
    if data.contact_line.strip():
        doc.add_paragraph(data.contact_line.strip())
    doc.add_paragraph(datetime.now().strftime("%B %d, %Y"))
    doc.add_paragraph(str(job.get("company") or "Hiring Team"))
    doc.add_paragraph(f"Re: {job.get('title') or 'Job Application'}")
    doc.add_paragraph("")
    for paragraph in data.cover_letter.splitlines():
        cleaned = paragraph.strip()
        if cleaned:
            doc.add_paragraph(cleaned)
    doc.save(path)


def _generate_documents(llm: Any, cv_text: str, profile: dict[str, Any], job: dict[str, Any], job_dir: Path) -> tuple[Path, Path]:
    result = llm.with_structured_output(TailoredApplication).invoke(_build_prompt(cv_text, profile, job))
    if not isinstance(result, TailoredApplication):
        result = TailoredApplication.model_validate(result)
    if not result.cover_letter.strip() or not result.summary.strip():
        raise ValueError("The AI response did not contain a usable summary and cover letter.")

    cv_path = job_dir / "tailored_cv.docx"
    letter_path = job_dir / "cover_letter.docx"
    _write_cv_docx(result, cv_path)
    _write_cover_letter_docx(result, job, letter_path)
    (job_dir / "application_metadata.json").write_text(
        json.dumps({
            "job_title": job.get("title", ""),
            "company": job.get("company", ""),
            "location": job.get("location", ""),
            "url": job.get("url", ""),
            "fit_score": job.get("fit_score", 0),
            "missing_skills": job.get("missing_skills", []),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "review_note": "Review every generated document before submitting. AI-generated content must be checked for accuracy."
        }, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return cv_path, letter_path


def _authorized_auto_submit(job: dict[str, Any], cv_path: Path, cover_letter_path: Path) -> dict[str, Any] | None:
    """Delegate to optional, explicitly implemented authorized adapters.

    Returns None when no adapter supports the platform. This starter package
    intentionally contains no scraper/browser automation for job-board forms.
    """
    if not AUTO_APPLY_ENABLED:
        return None
    try:
        from application_adapters import try_submit_with_authorized_adapter
        return try_submit_with_authorized_adapter(job, cv_path, cover_letter_path)
    except ImportError:
        logger.info("No application_adapters.py found; this job needs manual submission.")
        return None
    except Exception:
        logger.exception("Authorized application adapter failed for %s", job.get("url", ""))
        return {"status": "submission_failed", "notes": "An authorized adapter reported an error; review and submit manually."}


def _recipient_for(candidate_is_secondary: bool) -> list[str]:
    if candidate_is_secondary:
        raw = os.getenv("SECOND_CV_EMAIL", "").strip()
        if not raw:
            raw = ""
    else:
        raw = os.getenv("MY_CV_EMAILS", "").strip() or os.getenv("NOTIFICATION_EMAILS", "").strip()
    return list(dict.fromkeys(x.strip() for x in raw.split(",") if x.strip()))


def process_applications(
    *,
    cv_path: Path,
    cv_text: str,
    profile: dict[str, Any],
    assessed_jobs: list[dict[str, Any]],
    db_path: Path,
    output_dir: Path,
    llm: Any,
    candidate_is_secondary: bool,
    run_number: int,
) -> dict[str, Any]:
    """Prepare per-job documents, attempt authorized adapters, and email a report.

    Jobs are persisted in the existing jobs table. Only unprocessed qualifying
    jobs are packaged, and failed preparation is retried on a later run.
    """
    ensure_application_table(db_path)
    candidate_label = "Secondary candidate" if candidate_is_secondary else cv_path.stem.replace("_", " ")
    application_root = output_dir / "applications" / _safe_slug(cv_path.stem)
    application_root.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc).isoformat()

    with sqlite3.connect(db_path) as connection:
        rows = connection.execute(
            """SELECT j.url, j.title, j.company, j.location, j.region, j.description,
                      j.fit_score, j.reason, j.missing_skills, j.recommended,
                      j.verification_status, j.verification_reason,
                      a.status, a.report_sent
               FROM jobs j LEFT JOIN applications a ON a.url = j.url
               WHERE COALESCE(j.fit_score, 0) >= ?
                 AND COALESCE(j.verification_status, 'unverified') IN ('confirmed_job', 'likely_job')
               ORDER BY j.fit_score DESC, j.last_seen DESC""",
            (APPLICATION_MIN_FIT_SCORE,),
        ).fetchall()

    candidates: list[dict[str, Any]] = []
    for row in rows:
        (url, title, company, location, region, description, score, reason,
         missing, recommended, verification_status, verification_reason,
         status, report_sent) = row
        if verification_status not in {"confirmed_job", "likely_job"}:
            logger.info("Skipping application package for unverified lead: %s | %s", title, url)
            continue
        if _is_listing_page({"url": url, "title": title}):
            logger.info("Skipping search/listing page even though its score is high: %s | %s", title, url)
            # Prevent this old database row from being treated as recommended again.
            with sqlite3.connect(db_path) as connection:
                connection.execute("UPDATE jobs SET recommended = 0 WHERE url = ?", (url,))
            continue
        if status in {"ready_for_manual_application", "submitted", "awaiting_manual_action"}:
            continue
        try:
            missing_skills = json.loads(missing or "[]")
        except Exception:
            missing_skills = [x.strip() for x in str(missing or "").split(",") if x.strip()]
        candidates.append({
            "url": url, "title": title or "Untitled job", "company": company or "Unknown company",
            "location": location or "Not specified", "region": region or "", "description": description or "",
            "fit_score": int(score or 0), "reason": reason or "", "missing_skills": missing_skills,
            "recommended": int(score or 0) >= APPLICATION_MIN_FIT_SCORE,
            "verification_status": verification_status,
            "verification_reason": verification_reason or "",
            "platform": _host(url),
        })

    newly_reportable: list[dict[str, Any]] = []
    prepared_count = 0
    for job in candidates:
        if prepared_count >= MAX_APPLICATION_PACKAGES_PER_RUN:
            logger.info("Reached MAX_APPLICATION_PACKAGES_PER_RUN=%d; remaining jobs stay queued.", MAX_APPLICATION_PACKAGES_PER_RUN)
            break
        url = job["url"]
        if not url:
            continue
        job_hash = hashlib.sha1(url.encode('utf-8')).hexdigest()[:8]
        job_dir = application_root / f"{_safe_slug(job['company'])}_{_safe_slug(job['title'])}_{job_hash}"
        job_dir.mkdir(parents=True, exist_ok=True)
        try:
            cv_doc, letter_doc = _generate_documents(llm, cv_text, profile, job, job_dir)
            result = _authorized_auto_submit(job, cv_doc, letter_doc)
            if result and result.get("status") == "submitted":
                status = "submitted"
                notes = str(result.get("notes") or "Submission confirmation returned by configured authorized adapter.")
                confirmation = str(result.get("confirmation") or "")
                submitted_at = datetime.now(timezone.utc).isoformat()
            elif result and result.get("status") == "submission_failed":
                status = "awaiting_manual_action"
                notes = str(result.get("notes") or "Submission could not be confirmed; manual review required.")
                confirmation = ""
                submitted_at = ""
            else:
                status = "ready_for_manual_application"
                notes = "No configured authorized adapter supports this job board; documents are prepared for manual submission."
                confirmation = ""
                submitted_at = ""

            with sqlite3.connect(db_path) as connection:
                connection.execute(
                    """INSERT INTO applications
                    (url, title, company, location, fit_score, platform, status, cv_path,
                     cover_letter_path, notes, created_at, updated_at, submitted_at,
                     confirmation, report_sent)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
                    ON CONFLICT(url) DO UPDATE SET
                      title=excluded.title, company=excluded.company, location=excluded.location,
                      fit_score=excluded.fit_score, platform=excluded.platform, status=excluded.status,
                      cv_path=excluded.cv_path, cover_letter_path=excluded.cover_letter_path,
                      notes=excluded.notes, updated_at=excluded.updated_at,
                      submitted_at=excluded.submitted_at, confirmation=excluded.confirmation""",
                    (url, job["title"], job["company"], job["location"], job["fit_score"], job["platform"], status,
                     str(cv_doc.resolve()), str(letter_doc.resolve()), notes, now, now, submitted_at, confirmation),
                )
            job.update({"status": status, "notes": notes, "cv_path": str(cv_doc.resolve()),
                        "cover_letter_path": str(letter_doc.resolve()), "confirmation": confirmation})
            newly_reportable.append(job)
            prepared_count += 1
            logger.info("Application package prepared: %s | %s | %s", job["title"], job["company"], status)
        except Exception as exc:
            logger.exception("Could not prepare application package for %s", url)
            with sqlite3.connect(db_path) as connection:
                connection.execute(
                    """INSERT INTO applications
                    (url, title, company, location, fit_score, platform, status, notes, created_at, updated_at, report_sent)
                    VALUES (?, ?, ?, ?, ?, ?, 'preparation_failed', ?, ?, ?, 0)
                    ON CONFLICT(url) DO UPDATE SET status='preparation_failed', notes=excluded.notes, updated_at=excluded.updated_at""",
                    (url, job["title"], job["company"], job["location"], job["fit_score"], job["platform"],
                     f"Document preparation failed: {type(exc).__name__}: {exc}", now, now),
                )
            newly_reportable.append({**job, "status": "preparation_failed", "notes": f"Document preparation failed: {type(exc).__name__}: {exc}"})
            prepared_count += 1

    if not newly_reportable:
        logger.info("No new application packages to report for %s.", candidate_label)
        return {"application_packages_prepared": 0}

    # One ZIP attachment per candidate/run makes all generated files easy to retrieve.
    zip_path = output_dir / f"application_packages_{_safe_slug(cv_path.stem)}_run_{run_number}.zip"
    generated_files: list[Path] = []
    for job in newly_reportable:
        for key in ("cv_path", "cover_letter_path"):
            value = job.get(key)
            if value and Path(value).exists():
                generated_files.append(Path(value))
    if generated_files:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for file_path in generated_files:
                try:
                    archive.write(file_path, arcname=str(file_path.resolve().relative_to(output_dir.resolve())))
                except ValueError:
                    archive.write(file_path, arcname=file_path.name)

    recipients = _recipient_for(candidate_is_secondary)
    if not recipients:
        logger.warning("No application report recipients configured for %s.", candidate_label)
        return {"application_packages_prepared": prepared_count}

    submitted = [j for j in newly_reportable if j.get("status") == "submitted"]
    manual = [j for j in newly_reportable if j.get("status") in {"ready_for_manual_application", "awaiting_manual_action"}]
    failed = [j for j in newly_reportable if j.get("status") == "preparation_failed"]
    lines = [
        f"Application report for {candidate_label}",
        f"Run: {run_number}",
        f"Minimum match score: {APPLICATION_MIN_FIT_SCORE}/100",
        "",
        "A submitted status is used only when an authorized adapter returned a submission result. Opening a page or preparing files does not count as applying.",
        "",
        f"SUCCESSFULLY SUBMITTED ({len(submitted)})",
    ]
    if submitted:
        for job in submitted:
            lines.extend([f"- {job['title']} — {job['company']} | {job['fit_score']}/100", f"  {job['url']}", f"  Confirmation: {job.get('confirmation') or 'Adapter confirmed submission; no reference supplied.'}"])
    else:
        lines.append("- No applications were confirmed as submitted in this run.")
    lines.extend(["", f"READY FOR MANUAL APPLICATION ({len(manual)})"])
    if manual:
        for job in manual:
            lines.extend([f"- {job['title']} — {job['company']} | {job['location']} | {job['fit_score']}/100", f"  Apply: {job['url']}", f"  Tailored CV: {job.get('cv_path', 'See ZIP attachment')}", f"  Cover letter: {job.get('cover_letter_path', 'See ZIP attachment')}", f"  Note: {job.get('notes', '')}"])
    else:
        lines.append("- No manual applications in this report.")
    lines.extend(["", f"DOCUMENT PREPARATION NEEDS REVIEW ({len(failed)})"])
    if failed:
        for job in failed:
            lines.extend([f"- {job['title']} — {job['company']}", f"  {job.get('url', '')}", f"  {job.get('notes', '')}"])
    else:
        lines.append("- None.")
    lines.extend(["", "Review all AI-generated documents before submitting. Vacancy status may change, and a discovered link is not proof that a job is still open."])
    text_body = "\n".join(lines)

    html_lines = ["<html><body>", "<p><strong>Important:</strong> A submitted status is used only when an authorized adapter returned a submission result. Opening a page or preparing files does not count as applying.</p>", f"<h2>Successfully submitted ({len(submitted)})</h2><ul>"]
    for job in submitted:
        safe_title, safe_company, safe_url = escape(str(job["title"])), escape(str(job["company"])), escape(str(job["url"]), quote=True)
        html_lines.append(f"<li><strong>{safe_title}</strong> — {safe_company} ({job['fit_score']}/100) — <a href=\"{safe_url}\">job link</a></li>")
    html_lines.append("</ul><h2>Ready for manual application</h2><ul>")
    for job in manual:
        safe_title, safe_company = escape(str(job["title"])), escape(str(job["company"]))
        safe_url, safe_notes = escape(str(job["url"]), quote=True), escape(str(job.get("notes", "")))
        html_lines.append(f"<li><strong>{safe_title}</strong> — {safe_company} ({job['fit_score']}/100)<br/><a href=\"{safe_url}\">Apply on website</a><br/>Tailored CV and cover letter are in the attached ZIP (if included).<br/>{safe_notes}</li>")
    if not manual:
        html_lines.append("<li>No manual applications in this report.</li>")
    html_lines.append(f"</ul><h2>Preparation failures ({len(failed)})</h2><ul>")
    for job in failed:
        html_lines.append(f"<li>{escape(str(job['title']))} — {escape(str(job['company']))}: {escape(str(job.get('notes', '')))}</li>")
    if not failed:
        html_lines.append("<li>None.</li>")
    html_lines.append("</ul><p>Review all AI-generated documents before submitting. A discovered link is not proof that a vacancy is still open.</p></body></html>")

    try:
        send_application_report_email(
            subject=f"{candidate_label} — Application Report: {len(submitted)} submitted, {len(manual)} need action",
            text_body=text_body,
            html_body="".join(html_lines),
            recipients=recipients,
            attachment_path=zip_path if zip_path.exists() else None,
        )
        urls_to_mark = [j["url"] for j in newly_reportable]
        with sqlite3.connect(db_path) as connection:
            connection.executemany("UPDATE applications SET report_sent=1 WHERE url=?", [(u,) for u in urls_to_mark])
        logger.info("Application report sent to %s", ", ".join(recipients))
    except Exception:
        logger.exception("Application report email failed; report_sent remains unset so it can be retried.")

    return {"application_packages_prepared": prepared_count, "application_report_path": str(zip_path) if zip_path.exists() else ""}
